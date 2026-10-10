#!/usr/bin/env python3
"""Render docs/GAME_DESIGN_DOC.md as a Word document (docs/GAME_DESIGN_DOC.docx).

The markdown is canonical; the .docx is a generated copy for reading/sharing in Word.
Needs python-docx (pip install python-docx). Handles the subset the GDD uses: # headings,
paragraphs (with indented continuation lines), - / 1. lists, > quotes, | tables, ---,
and **bold** / *italic* / `code` inline.

    python3 tools/gdd_to_docx.py [out.docx]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "GAME_DESIGN_DOC.md"
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x7A, 0x2E, 0x2E)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def build(lines, doc):
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip("\n")
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 3) if level > 1 else 0)
            add_inline(h, m.group(2))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            head = split_row(rows[0])
            body = [split_row(r) for r in rows[2:]]
            t = doc.add_table(rows=1 + len(body), cols=len(head))
            t.style = "Table Grid"
            for c, txt in enumerate(head):
                cell = t.rows[0].cells[c]
                cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], txt)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
                shade(cell, "E4DCCB")
            for ri, row in enumerate(body, start=1):
                for c in range(len(head)):
                    txt = row[c] if c < len(row) else ""
                    cell = t.rows[ri].cells[c]
                    cell.paragraphs[0].text = ""
                    add_inline(cell.paragraphs[0], txt)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_inline(p, " ".join(b for b in buf if b))
            for r in p.runs:
                r.italic = True
            continue
        mb = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if mb:
            indent = len(mb.group(1))
            numbered = mb.group(2)[0].isdigit()
            text = mb.group(3)
            i += 1
            # indented continuation lines belong to this item
            while i < n and lines[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]) \
                    and not lines[i].strip().startswith(("|", "#", ">")):
                text += " " + lines[i].strip()
                i += 1
            style = "List Number" if numbered else "List Bullet"
            if indent >= 2:
                style += " 2"
            p = doc.add_paragraph(style=style)
            add_inline(p, text)
            continue
        # paragraph: gather following non-blank, non-structural lines
        text = s
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||>|---$|\s*([-*]|\d+\.)\s)", lines[i].strip()):
            text += " " + lines[i].strip()
            i += 1
        p = doc.add_paragraph()
        add_inline(p, text)


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "GAME_DESIGN_DOC.docx"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    build(SRC.read_text(encoding="utf-8").splitlines(), doc)
    doc.core_properties.title = "DF30 - Game Design Doc"
    doc.core_properties.author = "Mammoth Games"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
